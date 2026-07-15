"""Extract text from PPT files for presentation building."""
from pptx import Presentation
import json
import os

def extract_pptx(path):
    prs = Presentation(path)
    slides = []
    for i, slide in enumerate(prs.slides):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
        slides.append({"slide": i+1, "content": texts})
    return slides

print("=== AGGREGATE SLIDES ===")
agg = extract_pptx(r"C:\Users\BHS1RNG\Downloads\Aggregate-12-05-2026-dca-detailed-slides.pptx")
for s in agg:
    print(f"\n--- Slide {s['slide']} ---")
    for t in s['content']:
        print(t)

print("\n\n=== BRIEF SLIDES ===")
brief = extract_pptx(r"C:\Users\BHS1RNG\Downloads\Brief-DCA_Overview_Q2-20-5-26.pptx")
for s in brief:
    print(f"\n--- Slide {s['slide']} ---")
    for t in s['content']:
        print(t)

# Save JSON for later use
with open("ppt_content.json", "w") as f:
    json.dump({"aggregate": agg, "brief": brief}, f, indent=2)
print("\n\nSaved to ppt_content.json")
