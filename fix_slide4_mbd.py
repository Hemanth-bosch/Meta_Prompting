"""
Patch AIQ-212-Overview-For-MIDAS.pptx:
  - Rebuild slide 4 (MBD-UTA) to match slide 3 (DCA) layout exactly
  - Delete slide 5 (verbose version)

Slide 3 (DCA) layout reference:
  Title bar:        y=648000  h=388800   sz=304800 (24pt)
  Subtitle bar:     y=259200  h=388800   top strip
  "How →" label:    x=379825  y=1203590  fill=PURPLE  sz=177800 bold
  4 optimizer cards (left half):
    Card 1 (Seed):     x=426656  y=1261576  w=3097998  h=1574800
    Card 2 (Bootstrap):x=3711204 y=1261576  w=2750556  h=1574800
    Card 3 (MIPROv2):  x=3663537 y=2920196  w=2798219  h=1524000
    Card 4 (GEPA):     x=426655  y=2920196  w=3097999  h=1524000
  Metrics label:    x=426656  y=4586993  fill=PURPLE  sz=177800 bold
  Metric box:       x=453857  y=4758811  w=6007899  fill=F4D5FF
  Results label:    x=6667084 y=1244412  fill=00B050  sz=177800 bold
  Results box:      x=6667084 y=1483827  fill=F0F0F0  w=4302541 h=1028791
  Chart/picture:    x=6654754 y=2665347  w=4169040  h=2857266
"""

import sys
sys.path.insert(0, r'C:\Users\BHS1RNG\.conda\envs\new2\Lib\site-packages')

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree
import copy

SRC = r"C:\Users\BHS1RNG\Downloads\AIQ-212-Overview-For-MIDAS.pptx"
DST = r"C:\Users\BHS1RNG\Downloads\AIQ-212-Overview-For-MIDAS.pptx"

prs = Presentation(SRC)
print(f"Loaded: {len(prs.slides)} slides")

# ── Colors ────────────────────────────────────────────────────────────────────
BOSCH_BLUE  = RGBColor(0x00, 0x7B, 0xC0)
DARK_NAVY   = RGBColor(0x20, 0x38, 0x64)
GREEN       = RGBColor(0x00, 0xB0, 0x50)
GREEN2      = RGBColor(0x92, 0xD0, 0x50)   # bright green for ROLE label
PURPLE      = RGBColor(0x9E, 0x28, 0x96)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
BLACK       = RGBColor(0x00, 0x00, 0x00)
DARK_GRAY   = RGBColor(0x33, 0x33, 0x33)
LIGHT_GRAY  = RGBColor(0xF0, 0xF0, 0xF0)
ORANGE      = RGBColor(0xFF, 0x6B, 0x00)
TEAL        = RGBColor(0x00, 0x7D, 0xC3)
METRIC_BG   = RGBColor(0xF4, 0xD5, 0xFF)   # lavender — metric box bg (matches slide 3)
CARD_DARK   = RGBColor(0x1F, 0x39, 0x64)   # dark blue card fill (same as slide 3 cards)

# ── Shape helpers (all EMU inputs) ────────────────────────────────────────────
def solid_rect(slide, x, y, w, h, fill_rgb, line_rgb=None, radius_val=None):
    shp = slide.shapes.add_shape(1, Emu(x), Emu(y), Emu(w), Emu(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill_rgb
    shp.line.color.rgb = fill_rgb if line_rgb is None else line_rgb
    if radius_val:
        sp = shp._element
        pg = sp.find('.//' + qn('a:prstGeom'))
        if pg is not None:
            pg.set('prst', 'roundRect')
            av = pg.find(qn('a:avLst'))
            if av is None:
                av = etree.SubElement(pg, qn('a:avLst'))
            av.clear()
            gd = etree.SubElement(av, qn('a:gd'))
            gd.set('name', 'adj')
            gd.set('fmla', f'val {radius_val}')
    return shp

def add_tb(slide, x, y, w, h, lines, sizes, bolds, colors,
           align=PP_ALIGN.LEFT, wrap=True):
    """Multi-line textbox. lines/sizes/bolds/colors are parallel lists."""
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    for i, (txt, sz, bold, col) in enumerate(zip(lines, sizes, bolds, colors)):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = txt
        run.font.size = Pt(sz)
        run.font.bold = bold
        run.font.color.rgb = col
    return tb

def card(slide, x, y, w, h, title, role_text, bullets, card_fill=CARD_DARK):
    """Optimizer card matching slide 3 style exactly.
    Title: 14pt bold WHITE
    ROLE: label: 12pt bold GREEN2
    Bullets: 12pt normal WHITE
    """
    solid_rect(slide, x, y, w, h, card_fill)
    lines  = [title, role_text] + bullets
    sizes  = [14, 12] + [12]*len(bullets)
    bolds  = [True, True] + [False]*len(bullets)
    colors = [WHITE, GREEN2] + [WHITE]*len(bullets)
    add_tb(slide, x+60_000, y+60_000, w-120_000, h-120_000,
           lines, sizes, bolds, colors)

# ── Step 1: Clear slide 4 content (keep layout placeholders) ─────────────────
slide4 = prs.slides[3]   # 0-indexed

# Remove all non-placeholder shapes
to_remove = []
for sh in list(slide4.shapes):
    # keep title + subtitle placeholders (idx 0, 1)
    try:
        if sh.is_placeholder and sh.placeholder_format.idx in (0, 1):
            continue
    except: pass
    to_remove.append(sh._element)
for el in to_remove:
    el.getparent().remove(el)

# Update title placeholder text
for ph in slide4.placeholders:
    idx = ph.placeholder_format.idx
    if idx == 0:   # title
        ph.text = ""
        tf = ph.text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = "MBD-UTA — Automated Prompt Optimization Pipeline"
        run.font.size = Pt(24)
        run.font.bold = None   # inherit
    elif idx == 1:  # subtitle / top bar
        ph.text = "MiDAS Collaboration"

# ── Step 2: Build MBD-UTA content matching slide 3 layout ────────────────────

# "How →" label (purple badge — same position as slide 3)
tb_how = solid_rect(slide4, 379825, 1203590, 1003145, 224302, PURPLE)
tf = tb_how.text_frame
tf.word_wrap = False
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
run = p.add_run()
run.text = "How →"
run.font.size = Pt(14)
run.font.bold = True
run.font.color.rgb = WHITE

# ── 4 optimizer/step cards (exact same x,y,w,h as slide 3) ───────────────────

# Card 1 — Prompt 1: Context Extraction  (top-left, was "Seed Generation")
card(slide4,
     426656, 1261576, 3097998, 1574800,
     "1.  Context Extraction  ← LLM Call 1",
     "ROLE: Understand the weakspot equation graph",
     [
         "Input: filtered equation JSON + signalsheet",
         "Identifies: alterable inputs, checkable outputs,",
         "  branch scenarios, boundary conditions",
         "Output: plain-text context analysis",
     ])

# Card 2 — Prompt 2: Test Case Generation  (top-right, was "Bootstrap")
card(slide4,
     3711204, 1261576, 2750556, 1574800,
     "2.  Test Case Generation  ← LLM Call 2",
     "ROLE: Generate structured JSON test cases",
     [
         "Input: context + equation JSON",
         "Output: Pre-Condition / Actions / Expected",
         "GEPA evolves this prompt for signal naming,",
         "timing rules, branch coverage",
     ])

# Card 3 — Post-processing  (bottom-left, was "GEPA")
card(slide4,
     426655, 2920196, 3097999, 1524000,
     "3.  Equation Solving + DLL Verify  ← LLM Call 4",
     "ROLE: Derive concrete input values for each test step",
     [
         "LLM proposes Variable / Parameter values",
         "DLL executes equation chain & verifies output",
         "Retry once on mismatch (DLL feedback to LLM)",
     ],
     card_fill=RGBColor(0x00, 0x45, 0x82))   # slightly different blue, like slide 3 variation

# Card 4 — GEPA optimization  (bottom-right, was "MIPROv2")
card(slide4,
     3663537, 2920196, 2798219, 1524000,
     "4.  GEPA Optimization",
     "ROLE: Reflective prompt rewriting",
     [
         "Optimizer: GEPA (auto=light)",
         "Reflection LM: gpt-5.2",
         "Optimizes Prompt 1 + Prompt 2",
         "Next: add Prompt 4 (Eq. Solving)",
     ],
     card_fill=CARD_DARK)

# ── Metrics label + box (matches slide 3 exactly) ────────────────────────────
tb_m = solid_rect(slide4, 426656, 4586993, 769846, 247654, PURPLE)
tf = tb_m.text_frame
tf.word_wrap = False
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
run = p.add_run()
run.text = "Metrics"
run.font.size = Pt(14)
run.font.bold = True
run.font.color.rgb = WHITE

metric_box = solid_rect(slide4, 453857, 4758811, 6007899, 737320, METRIC_BG)
add_tb(slide4, 513857, 4798811, 5907899, 660000,
       [
           "Score  =  25%×M1(I/O)  +  20%×M2(Signals)  +  10%×M3(Structure)  +  25%×M4(Branches)  +  5%×M6(Bounds)  +  15%×M7(DLL verify)",
           "Reference-free: no ground truth needed — scored from equation graph + signalsheet + DLL",
       ],
       [12, 11], [True, False], [BLACK, DARK_GRAY])

# ── Results label (green badge, matches slide 3) ─────────────────────────────
tb_res = solid_rect(slide4, 6667084, 1244412, 1494422, 282829, GREEN)
tf = tb_res.text_frame
tf.word_wrap = False
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
run = p.add_run()
run.text = "Results"
run.font.size = Pt(14)
run.font.bold = True
run.font.color.rgb = WHITE

# ── Results content box (matches slide 3 size/position) ──────────────────────
res_box = solid_rect(slide4, 6667084, 1483827, 4302541, 1028791, LIGHT_GRAY)
add_tb(slide4, 6727084, 1513827, 4182541, 970000,
       [
           "Baseline (seed prompts):   73.9%",
           "GEPA Best (iter 7):           91.0%   +17.1 pp",
           "Pareto Best (iter 13):        92.1%",
           "",
           "Run stopped at 27%  (198 / 742 rollouts)",
           "M7 DLL = stub 0.5 throughout  (15% blind)",
           "Prompts 3+4 not yet optimized in this run",
       ],
       [12, 12, 12, 9, 11, 10, 10],
       [True, True, True, False, False, False, False],
       [DARK_NAVY, GREEN, TEAL, WHITE, DARK_GRAY, ORANGE, DARK_GRAY])

# ── Next steps box (right side, below results — fills the chart area) ─────────
# In slide 3 this area has a chart/picture. Use a styled box instead.
add_tb(slide4, 6667084, 2560000, 4302541, 200000,
       ["Next Steps"],
       [12], [True], [DARK_NAVY])

next_box = solid_rect(slide4, 6667084, 2720000, 4302541, 2800000,
                      RGBColor(0xE8, 0xF4, 0xFF))
add_tb(slide4, 6727084, 2760000, 4182541, 2700000,
       [
           "Full pipeline: all 3 prompts optimizable with GEPA",
           "  Prompt 1  Context Extraction      ← already optimized",
           "  Prompt 2  Test Case Generator     ← already optimized",
           "  Prompt 4  Equation Solving DLL   ← next to add",
           "",
           "Real M7: DLL executes test cases, verifies outputs",
           "Staged metric: M1/M3/M4/M6 → Prompt 2 feedback",
           "                      M2/M7        → Prompt 4 feedback",
           "",
           "True MIDAS baseline → then full GEPA run",
           "Scale to full test set evaluation",
       ],
       [11, 10, 10, 10, 9, 11, 10, 10, 9, 11, 10],
       [True, False, False, False, False, True, False, False, False, True, False],
       [DARK_NAVY, DARK_GRAY, GREEN, BOSCH_BLUE, WHITE,
        DARK_NAVY, DARK_GRAY, DARK_GRAY, WHITE, DARK_NAVY, DARK_GRAY])

# ── Step 3: Delete slide 5 ────────────────────────────────────────────────────
# slide index 4 = slide 5 (0-based)
rId = prs.slides._sldIdLst[4].rId
prs.part.drop_rel(rId)
del prs.slides._sldIdLst[4]

print(f"Slides after patch: {len(prs.slides)}")

prs.save(DST)
print(f"Saved: {DST}")
